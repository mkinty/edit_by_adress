"""Application de la modification par adresse à un classeur Excel.

Le choix des lignes n'est plus dicté par un ID erreur mais par une
**adresse** : une ligne n'est retenue que si la valeur de sa colonne
« Adresse » correspond à l'une des adresses saisies. Sur chaque ligne
retenue, l'unique opération possible est d'écrire dans la colonne « ID
erreur » la valeur associée à cette adresse dans la saisie.

Une même adresse peut être saisie plusieurs fois avec des valeurs
différentes — plusieurs logements à la même adresse, par exemple, chacun
avec son propre ID erreur à corriger. Dans ce cas, les lignes du fichier
partageant cette adresse reçoivent les valeurs saisies **dans l'ordre** :
la 1re ligne rencontrée (de haut en bas) reçoit la 1re valeur saisie pour
cette adresse, la 2e ligne la 2e valeur, etc. (cf. ``_valeur_suivante``).

Le traitement se fait en **deux temps**, pour ne payer le prix fort que
lorsqu'il y a réellement quelque chose à écrire :

1. **Repérage** — le XML de la feuille est parcouru en flux sur la seule
   colonne « Adresse » (cf. ``services.lecture_rapide``, avec repli sur
   openpyxl si le classeur a une structure inattendue). Cette passe ne
   construit aucun objet cellule et détermine, pour chaque ligne concernée,
   la valeur à écrire dans la colonne « ID erreur ». Les deux colonnes sont
   résolues ici : une colonne absente échoue donc avant tout chargement
   coûteux.
2. **Écriture** — le classeur n'est chargé entièrement que si au moins une
   ligne correspond.

Un fichier dont aucune adresse ne correspond n'est ainsi jamais chargé en
mémoire ; c'est le gain le plus important sur les gros classeurs.

Attention : openpyxl réécrit intégralement le classeur à la sauvegarde.
Certains éléments (graphiques, images, mises en forme conditionnelles
avancées) peuvent ne pas être préservés. Testez toujours sur une copie avant
un traitement en masse. Le fichier n'est réécrit que si au moins une cellule
a réellement changé.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field

from openpyxl import load_workbook
import warnings

from ..models.config_adresses import ConfigAdresses
from .adresses import SaisieAdresse, normaliser_adresse
from .colonnes_excel import resoudre_colonne, table_depuis_entetes
from .lecture_rapide import LectureImpossible, lire_colonne, lire_entetes
from .valeurs import harmoniser_forme, memes_valeurs, valeur_a_ecrire

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Signature du rapporteur d'étapes (journal en temps réel).
Etape = Callable[[str], None]

# Périodicité des points d'avancement pendant le repérage des lignes.
_PAS_AVANCEMENT = 20_000

# Nombre de lignes explorées au maximum pour trouver une cellule de référence
# de mise en forme.
_PROFONDEUR_MODELE = 500

# Étendue de l'échantillon servant à déterminer la police dominante du classeur.
_LIGNES_POLICE = 200
_COLONNES_POLICE = 40


class ColonneIntrouvable(LookupError):
    """La colonne des adresses ou celle des ID erreur est absente du classeur."""


@dataclass
class BilanFichier:
    """Ce qu'a produit le traitement d'un fichier.

    ``adresses_trouvees`` permet à l'appelant de signaler les adresses
    demandées qui n'apparaissent nulle part dans le classeur.
    """

    nb_modifications: int = 0
    nb_lignes: int = 0
    adresses_trouvees: set[str] = field(default_factory=set)
    nb_lignes_lues: int = 0
    enregistre: bool = False


def appliquer_adresses(
    chemin_fichier: str,
    config: ConfigAdresses,
    saisies: Iterable[SaisieAdresse],
    simuler: bool = False,
    on_etape: Etape | None = None,
) -> BilanFichier:
    """Écrit, pour chaque adresse trouvée, sa valeur d'ID erreur.

    Args:
        chemin_fichier: classeur à modifier (modifié sur place).
        config: colonnes et ligne d'en-tête du classeur.
        saisies: paires (adresse, valeur d'ID erreur) à traiter dans ce
            fichier — normalement toutes rattachées à la même commune.
        simuler: si vrai, rien n'est écrit sur le disque — le décompte
            retourné indique ce qui *serait* modifié.
        on_etape: appelé à chaque étape avec un message court, pour alimenter
            le journal pendant le traitement plutôt qu'à la fin.

    Returns:
        Un :class:`BilanFichier` : nombre de modifications, nombre de lignes
        retenues, et adresses effectivement rencontrées.

    Raises:
        ColonneIntrouvable: si la colonne des adresses ou celle des ID erreur
            est absente du classeur.
    """
    # Une adresse peut être saisie plusieurs fois avec des valeurs différentes
    # (plusieurs lignes distinctes du fichier partagent cette adresse) : les
    # valeurs sont mises en file et consommées dans l'ordre des lignes du
    # fichier — cf. _valeur_suivante.
    attendues: dict[str, list[str]] = {}
    for saisie in saisies:
        cle = normaliser_adresse(saisie.adresse)
        if cle:
            attendues.setdefault(cle, []).append(saisie.valeur_id)

    bilan = BilanFichier()
    if not attendues:
        return bilan

    _dire(on_etape, "analyse du fichier…")
    lignes, valeurs_par_ligne, index_id = _reperer(chemin_fichier, config, attendues, bilan, on_etape)

    if not lignes:
        _dire(on_etape, f"aucune ligne concernée sur {bilan.nb_lignes_lues} lue(s) — non ouvert en écriture")
        return bilan

    _dire(on_etape, f"{len(lignes)} ligne(s) à modifier — chargement du classeur…")
    classeur = load_workbook(chemin_fichier)
    try:
        feuille = classeur.active
        modele = _cellule_modele(feuille, index_id, config.ligne_entete, set(lignes))
        police = None if modele is not None else _police_dominante(feuille, config.ligne_entete)

        _dire(on_etape, "application des modifications…")
        for ligne in lignes:
            valeur = valeur_a_ecrire(valeurs_par_ligne[ligne])
            cellule = feuille.cell(ligne, index_id)
            if memes_valeurs(cellule.value, valeur):
                continue
            cellule.value = valeur
            if config.harmoniser_style:
                harmoniser_forme(cellule, modele, police)
            bilan.nb_modifications += 1

        if bilan.nb_modifications and not simuler:
            _dire(on_etape, f"enregistrement de {bilan.nb_modifications} modification(s)…")
            classeur.save(chemin_fichier)
            bilan.enregistre = True
        elif simuler:
            _dire(on_etape, "simulation — aucune écriture")
        else:
            _dire(on_etape, "valeurs déjà en place — fichier inchangé")
        return bilan
    finally:
        classeur.close()


# ── Passe 1 : repérage ────────────────────────────────────────────
def _reperer(
    chemin: str,
    config: ConfigAdresses,
    attendues: dict[str, list[str]],
    bilan: BilanFichier,
    on_etape: Etape | None,
) -> tuple[list[int], dict[int, str], int]:
    """Numéros de lignes concernées, valeur à écrire par ligne, colonne cible.

    La lecture directe du XML est tentée d'abord ; tout classeur de structure
    inattendue bascule sur openpyxl, plus lent mais universel. Le résultat est
    identique dans les deux cas.
    """
    try:
        return _reperer_rapide(chemin, config, attendues, bilan, on_etape)
    except LectureImpossible as exc:
        _dire(on_etape, f"lecture rapide indisponible ({exc}) — repli openpyxl")
        return _reperer_openpyxl(chemin, config, attendues, bilan, on_etape)


def _valeur_suivante(attendues: dict[str, list[str]], compteurs: dict[str, int], cle: str) -> str:
    """Valeur à écrire pour la prochaine ligne portant l'adresse ``cle``.

    Plusieurs lignes du fichier peuvent partager la même adresse (plusieurs
    logements à une même adresse, par exemple) tout en attendant des ID
    erreur différents : les valeurs saisies pour cette adresse sont donc
    consommées dans l'ordre, une par ligne rencontrée. S'il y a plus de
    lignes que de valeurs saisies, la dernière valeur est réutilisée pour les
    lignes excédentaires plutôt que de les laisser sans écriture.
    """
    valeurs = attendues[cle]
    index = compteurs.get(cle, 0)
    compteurs[cle] = index + 1
    return valeurs[min(index, len(valeurs) - 1)]


def _reperer_rapide(
    chemin: str, config: ConfigAdresses, attendues: dict[str, list[str]],
    bilan: BilanFichier, on_etape: Etape | None,
) -> tuple[list[int], dict[int, str], int]:
    """Repérage par lecture directe du XML : une seule colonne est décodée."""
    entetes = lire_entetes(chemin, config.ligne_entete)
    if not entetes:
        raise LectureImpossible("ligne d'en-tête vide")

    colonne_adresse, colonne_id = _resoudre_colonnes(entetes, config)

    lignes: list[int] = []
    valeurs_par_ligne: dict[int, str] = {}
    trouvees: set[str] = set()
    compteurs: dict[str, int] = {}
    derniere = config.ligne_entete
    for ligne, valeur in lire_colonne(chemin, colonne_adresse, config.ligne_entete + 1):
        derniere = ligne
        cle = normaliser_adresse(valeur)
        if cle in attendues:
            lignes.append(ligne)
            valeurs_par_ligne[ligne] = _valeur_suivante(attendues, compteurs, cle)
            trouvees.add(cle)
        if ligne % _PAS_AVANCEMENT == 0:
            _dire(on_etape, _avancement(ligne))

    return _cloturer(bilan, lignes, trouvees, derniere, config), valeurs_par_ligne, colonne_id


def _reperer_openpyxl(
    chemin: str, config: ConfigAdresses, attendues: dict[str, list[str]],
    bilan: BilanFichier, on_etape: Etape | None,
) -> tuple[list[int], dict[int, str], int]:
    """Repérage de repli, en lecture seule openpyxl."""
    from .colonnes_excel import entetes_depuis_feuille

    classeur = load_workbook(chemin, read_only=True)
    try:
        feuille = classeur.active
        entetes = entetes_depuis_feuille(feuille, config.ligne_entete)
        colonne_adresse, colonne_id = _resoudre_colonnes(entetes, config)

        lignes: list[int] = []
        valeurs_par_ligne: dict[int, str] = {}
        trouvees: set[str] = set()
        compteurs: dict[str, int] = {}
        ligne = config.ligne_entete
        for cellules in feuille.iter_rows(
            min_row=config.ligne_entete + 1,
            min_col=colonne_adresse, max_col=colonne_adresse,
            values_only=True,
        ):
            ligne += 1
            cle = normaliser_adresse(cellules[0] if cellules else None)
            if cle in attendues:
                lignes.append(ligne)
                valeurs_par_ligne[ligne] = _valeur_suivante(attendues, compteurs, cle)
                trouvees.add(cle)
            if ligne % _PAS_AVANCEMENT == 0:
                _dire(on_etape, _avancement(ligne))

        return _cloturer(bilan, lignes, trouvees, ligne, config), valeurs_par_ligne, colonne_id
    finally:
        classeur.close()


def _resoudre_colonnes(entetes: dict[int, str], config: ConfigAdresses) -> tuple[int, int]:
    """Colonne des adresses et colonne des ID erreur, résolues avant tout chargement.

    Faire cette résolution ici garantit qu'une colonne absente est signalée
    sans avoir chargé le classeur en entier.
    """
    table = table_depuis_entetes(entetes)
    colonne_adresse = _resoudre(config.colonne_adresse, table, "colonne des adresses")
    colonne_id = _resoudre(config.colonne_id, table, "colonne des ID erreur")
    return colonne_adresse, colonne_id


def _cloturer(
    bilan: BilanFichier, lignes: list[int], trouvees: set[str],
    derniere: int, config: ConfigAdresses,
) -> list[int]:
    """Reporte le résultat du repérage dans le bilan (une fois la passe finie)."""
    bilan.adresses_trouvees = trouvees
    bilan.nb_lignes = len(lignes)
    bilan.nb_lignes_lues = max(derniere - config.ligne_entete, 0)
    return lignes


def _avancement(ligne: int) -> str:
    return f"{ligne:,} lignes analysées…".replace(",", " ")


# ── Passe 2 : mise en forme reprise de la colonne écrite ──────────
def _cellule_modele(feuille, index: int, ligne_entete: int, exclues: set[int]):
    """Première cellule renseignée de la colonne, hors lignes modifiées.

    Le modèle est cherché parmi les lignes **non modifiées** : on recopie
    ainsi la forme d'origine du fichier, jamais celle d'une cellule qu'on
    vient d'écrire. L'exploration est bornée : au-delà de quelques centaines
    de lignes sans aucune valeur, la colonne est considérée comme vide.
    """
    derniere = min(feuille.max_row or ligne_entete, ligne_entete + _PROFONDEUR_MODELE)
    for ligne in range(ligne_entete + 1, derniere + 1):
        if ligne in exclues:
            continue
        cellule = _cellule_existante(feuille, ligne, index)
        if cellule is not None and cellule.value not in (None, ""):
            return cellule
    return None


def _police_dominante(feuille, ligne_entete: int):
    """Police la plus employée dans les données du classeur.

    Sert de repli quand la colonne cible est **entièrement vide** — cas
    normal d'une colonne d'ID erreur qu'on remplit précisément parce qu'elle
    ne l'est pas. Beaucoup d'audits sont composés en Calibri 10 alors
    qu'openpyxl écrit en Calibri 11 par défaut : sans ce relais, la valeur
    inscrite se repérerait immédiatement à l'œil.
    """
    derniere = min(feuille.max_row or ligne_entete, ligne_entete + _LIGNES_POLICE)
    colonnes = min(feuille.max_column or 1, _COLONNES_POLICE)

    frequences: Counter = Counter()
    exemples: dict = {}
    for ligne in range(ligne_entete + 1, derniere + 1):
        for colonne in range(1, colonnes + 1):
            cellule = _cellule_existante(feuille, ligne, colonne)
            if cellule is None or cellule.value in (None, ""):
                continue
            signature = (cellule.font.name, cellule.font.size)
            frequences[signature] += 1
            exemples.setdefault(signature, cellule.font)

    if not frequences:
        return None
    return exemples[frequences.most_common(1)[0][0]]


def _cellule_existante(feuille, ligne: int, colonne: int):
    """Cellule déjà présente dans la feuille, sans en créer une nouvelle.

    ``feuille.cell(...)`` crée la cellule — et donc sa ligne — dès qu'on la
    consulte : inspecter la mise en forme ajouterait ainsi des lignes vides au
    classeur enregistré. On interroge donc directement la table interne, en se
    rabattant sur l'accès normal si openpyxl venait à changer.
    """
    cellules = getattr(feuille, "_cells", None)
    if cellules is None:  # pragma: no cover — filet de sécurité
        return feuille.cell(ligne, colonne)
    return cellules.get((ligne, colonne))


# ── Helpers ───────────────────────────────────────────────────────
def _resoudre(reference: str, table: dict[str, int], contexte: str) -> int:
    """Traduit une référence de colonne en numéro, ou échoue explicitement."""
    index = resoudre_colonne(reference, table)
    if index is None:
        raise ColonneIntrouvable(
            f"{contexte} : colonne « {reference} » introuvable dans le fichier"
        )
    return index


def _dire(on_etape: Etape | None, message: str) -> None:
    if on_etape is not None:
        on_etape(message)
