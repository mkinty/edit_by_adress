"""Fixtures partagées : création de classeurs de test."""

import pytest
from openpyxl import Workbook

# En-tête utilisé par la plupart des tests : la colonne C porte les ID erreur.
ENTETES = {
    "A": "Adresse",
    "B": "Audit - Nb BAL",
    "C": "ID erreur",
    "Z": "Audit - Street View",
}


@pytest.fixture
def creer_classeur(tmp_path):
    """Retourne une fonction qui écrit un classeur et renvoie son chemin.

    ``colonnes`` associe une lettre de colonne à la liste des valeurs des
    lignes de données (à partir de la ligne 2).
    """

    def _creer(colonnes: dict[str, list], nom: str = "audit_test.xlsx", entetes=ENTETES):
        wb = Workbook()
        ws = wb.active
        for lettre, libelle in entetes.items():
            ws[f"{lettre}1"] = libelle
        for lettre, valeurs in colonnes.items():
            for i, valeur in enumerate(valeurs, start=2):
                ws[f"{lettre}{i}"] = valeur
        chemin = str(tmp_path / nom)
        wb.save(chemin)
        return chemin

    return _creer
