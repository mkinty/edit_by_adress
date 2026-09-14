"""Application de la modification, filtrée par adresse."""

import os

import pytest
from openpyxl import load_workbook

from sna.models.config_adresses import ConfigAdresses
from sna.services.adresses import SaisieAdresse
from sna.services.excel import ColonneIntrouvable, appliquer_adresses

VALEUR = "13001_9"


def _config(ligne_entete=1, colonne_adresse="Adresse", colonne_id="ID erreur",
            harmoniser_style=True):
    return ConfigAdresses(
        colonne_adresse=colonne_adresse, colonne_id=colonne_id,
        ligne_entete=ligne_entete, harmoniser_style=harmoniser_style,
    )


def _saisie(adresse, valeur=VALEUR):
    return SaisieAdresse(adresse, valeur)


# ── Filtrage par adresse ──────────────────────────────────────────
def test_seules_les_lignes_des_adresses_demandees_sont_modifiees(creer_classeur):
    chemin = creer_classeur({
        "A": ["12 rue A", "20 rue B", "5 rue C"],
        "C": ["x", "x", "x"],
    })

    bilan = appliquer_adresses(
        chemin, _config(), [_saisie("12 rue A", "13001_1"), _saisie("5 rue C", "13001_3")]
    )

    assert bilan.nb_modifications == 2
    assert bilan.nb_lignes == 2
    ws = load_workbook(chemin).active
    assert ws["C2"].value == "13001_1"
    assert ws["C3"].value == "x"        # adresse non demandée : ligne intacte
    assert ws["C4"].value == "13001_3"


def test_adresses_trouvees_et_absentes_sont_remontees(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    bilan = appliquer_adresses(
        chemin, _config(), [_saisie("12 rue A"), _saisie("adresse fantôme")]
    )

    assert bilan.adresses_trouvees == {"12 RUE A"}


def test_correspondance_tolerante_a_la_casse_et_aux_espaces(creer_classeur):
    chemin = creer_classeur({"A": [" 12   rue de la paix "], "C": ["x"]})

    bilan = appliquer_adresses(chemin, _config(), [_saisie("12 RUE DE LA PAIX")])

    assert bilan.nb_modifications == 1


def test_aucune_adresse_demandee_ne_modifie_rien(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    bilan = appliquer_adresses(chemin, _config(), [])

    assert bilan.nb_modifications == 0
    assert load_workbook(chemin).active["C2"].value == "x"


def test_l_entete_n_est_jamais_modifie(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    appliquer_adresses(chemin, _config(), [_saisie("12 rue A")])

    ws = load_workbook(chemin).active
    assert ws["C1"].value == "ID erreur"


def test_plusieurs_lignes_portant_la_meme_adresse(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A", "12 rue A"], "C": ["a", "b"]})

    bilan = appliquer_adresses(chemin, _config(), [_saisie("12 rue A")])

    assert bilan.nb_modifications == 2 and bilan.nb_lignes == 2
    ws = load_workbook(chemin).active
    assert ws["C2"].value == VALEUR and ws["C3"].value == VALEUR


def test_premiere_valeur_appliquee_quand_une_seule_ligne_correspond(creer_classeur):
    """Deux valeurs saisies pour une adresse qui n'a qu'une ligne : la 1re est retenue."""
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    appliquer_adresses(
        chemin, _config(), [_saisie("12 rue A", "premier"), _saisie("12 rue A", "second")]
    )

    assert load_workbook(chemin).active["C2"].value == "premier"


def test_valeurs_distribuees_dans_l_ordre_sur_les_lignes_dupliquees(creer_classeur):
    """Deux logements à la même adresse : chaque ligne reçoit son propre ID erreur."""
    chemin = creer_classeur({"A": ["12 rue A", "12 rue A", "12 rue A"], "C": ["x", "y", "z"]})

    bilan = appliquer_adresses(
        chemin, _config(), [_saisie("12 rue A", "premier"), _saisie("12 rue A", "second")]
    )

    assert bilan.nb_modifications == 3
    ws = load_workbook(chemin).active
    assert ws["C2"].value == "premier"
    assert ws["C3"].value == "second"
    assert ws["C4"].value == "second"   # valeur excédentaire : la dernière est réutilisée


# ── Désignation des colonnes ──────────────────────────────────────
def test_colonne_adresse_designee_par_sa_lettre(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    bilan = appliquer_adresses(
        chemin, _config(colonne_adresse="A"), [_saisie("12 rue A")]
    )

    assert bilan.nb_modifications == 1


def test_colonne_id_designee_par_sa_lettre(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    bilan = appliquer_adresses(chemin, _config(colonne_id="C"), [_saisie("12 rue A")])

    assert bilan.nb_modifications == 1


def test_nom_de_colonne_tolerant_aux_espaces_et_a_la_casse(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})
    config = _config(colonne_adresse="adresse", colonne_id="iderreur")

    assert appliquer_adresses(chemin, config, [_saisie("12 rue A")]).nb_modifications == 1


def test_colonne_id_inconnue_leve_une_erreur(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    with pytest.raises(ColonneIntrouvable) as erreur:
        appliquer_adresses(
            chemin, _config(colonne_id="Absente"), [_saisie("12 rue A")]
        )
    assert "ID erreur" in str(erreur.value)


def test_colonne_adresse_absente_leve_une_erreur(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    with pytest.raises(ColonneIntrouvable) as erreur:
        appliquer_adresses(
            chemin, _config(colonne_adresse="Absente"), [_saisie("12 rue A")]
        )
    assert "adresses" in str(erreur.value)


# ── Valeurs écrites ───────────────────────────────────────────────
def test_valeur_numerique_ecrite_comme_nombre(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "B": ["à chiffrer"]})

    appliquer_adresses(
        chemin, _config(colonne_id="B"), [_saisie("12 rue A", "42")]
    )

    assert load_workbook(chemin).active["B2"].value == 42


def test_valeur_ressemblant_a_un_id_reste_du_texte(creer_classeur):
    """Écrire « 13001_2 » ne doit pas produire le nombre 130012."""
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    appliquer_adresses(chemin, _config(), [_saisie("12 rue A", "13001_2")])

    assert load_workbook(chemin).active["C2"].value == "13001_2"


def test_valeur_deja_presente_n_est_pas_recomptee(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": [VALEUR]})

    assert appliquer_adresses(chemin, _config(), [_saisie("12 rue A")]).nb_modifications == 0


# ── Écriture sur disque ───────────────────────────────────────────
def test_aucune_modification_ne_reecrit_pas_le_fichier(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": [VALEUR]})
    mtime_avant = os.path.getmtime(chemin)

    appliquer_adresses(chemin, _config(), [_saisie("12 rue A")])

    assert os.path.getmtime(chemin) == mtime_avant


def test_simulation_ne_modifie_pas_le_fichier(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})

    bilan = appliquer_adresses(chemin, _config(), [_saisie("12 rue A")], simuler=True)

    assert bilan.nb_modifications == 1
    assert load_workbook(chemin).active["C2"].value == "x"


# ── Ligne d'en-tête ───────────────────────────────────────────────
def test_ligne_entete_personnalisee(creer_classeur):
    chemin = creer_classeur(
        {"A": ["Adresse", "12 rue A"], "B": ["ID erreur", "x"]}, entetes={}
    )
    config = _config(ligne_entete=2, colonne_adresse="Adresse", colonne_id="ID erreur")

    bilan = appliquer_adresses(chemin, config, [_saisie("12 rue A")])

    assert bilan.nb_modifications == 1
    assert load_workbook(chemin).active["B3"].value == VALEUR


# ── Repli sur openpyxl ────────────────────────────────────────────
def test_repli_openpyxl_donne_le_meme_resultat(creer_classeur, monkeypatch):
    """Un classeur illisible par le lecteur rapide reste traité correctement."""
    from sna.services import excel as module_excel
    from sna.services.lecture_rapide import LectureImpossible

    def refuser(*_args, **_kwargs):
        raise LectureImpossible("structure simulée")

    chemin = creer_classeur({
        "A": ["12 rue A", "20 rue B", "5 rue C"],
        "C": ["a", "b", "c"],
    })
    monkeypatch.setattr(module_excel, "lire_entetes", refuser)
    messages = []

    bilan = appliquer_adresses(
        chemin, _config(), [_saisie("12 rue A", "v1"), _saisie("5 rue C", "v3")],
        on_etape=messages.append,
    )

    assert bilan.nb_modifications == 2
    assert bilan.adresses_trouvees == {"12 RUE A", "5 RUE C"}
    assert any("repli openpyxl" in m for m in messages)
    ws = load_workbook(chemin).active
    assert [ws[f"C{ligne}"].value for ligne in (2, 3, 4)] == ["v1", "b", "v3"]


# ── Étapes journalisées ───────────────────────────────────────────
def test_les_etapes_sont_rapportees_au_fil_du_traitement(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})
    messages = []

    appliquer_adresses(chemin, _config(), [_saisie("12 rue A")], on_etape=messages.append)

    assert messages[0].startswith("analyse")
    assert any("chargement" in m for m in messages)
    assert any("enregistrement" in m for m in messages)


def test_fichier_sans_correspondance_n_est_pas_charge(creer_classeur):
    """Le message doit dire explicitement que le classeur n'a pas été ouvert."""
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})
    messages = []

    appliquer_adresses(
        chemin, _config(), [_saisie("adresse fantôme")], on_etape=messages.append
    )

    assert any("non ouvert en écriture" in m for m in messages)
    assert not any("chargement" in m for m in messages)


def test_colonne_absente_detectee_avant_tout_chargement(creer_classeur):
    chemin = creer_classeur({"A": ["12 rue A"], "C": ["x"]})
    messages = []

    with pytest.raises(ColonneIntrouvable):
        appliquer_adresses(
            chemin, _config(colonne_id="Absente"), [_saisie("12 rue A")],
            on_etape=messages.append,
        )

    assert not any("chargement" in m for m in messages)
