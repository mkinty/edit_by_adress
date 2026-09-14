"""Lecture rapide du XML d'un classeur.

Les classeurs rencontrés en production ne se ressemblent pas : Excel range les
textes dans une table partagée, openpyxl les écrit « en ligne », et l'onglet
actif n'est pas toujours le premier. Chaque cas est vérifié ici, ainsi que le
refus explicite des structures que le lecteur ne sait pas interpréter — c'est
ce refus qui déclenche le repli sur openpyxl.
"""

import zipfile

import pytest
from openpyxl import Workbook

from sna.services import lecture_rapide
from sna.services.lecture_rapide import (
    LectureImpossible, lire_colonne, lire_entetes,
)

# ── Classeurs fabriqués à la main (style Excel : chaînes partagées) ──
_CONTENT_TYPES = """<?xml version="1.0"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="xml" ContentType="application/xml"/>
</Types>"""

_RELS = """<?xml version="1.0"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Target="xl/workbook.xml"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"/>
</Relationships>"""

_WORKBOOK_RELS = """<?xml version="1.0"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Target="worksheets/sheet1.xml"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/>
  <Relationship Id="rId2" Target="worksheets/sheet2.xml"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/>
</Relationships>"""

_NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
_NS_REL = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def _workbook_xml(onglet_actif: int = 0) -> str:
    return f"""<?xml version="1.0"?>
<workbook {_NS} {_NS_REL}>
  <bookViews><workbookView activeTab="{onglet_actif}"/></bookViews>
  <sheets>
    <sheet name="Premier" sheetId="1" r:id="rId1"/>
    <sheet name="Second" sheetId="2" r:id="rId2"/>
  </sheets>
</workbook>"""


def _feuille_xml(lignes) -> bytes:
    if isinstance(lignes, str):
        lignes = lignes.encode("utf-8")
    entete = f'<?xml version="1.0"?><worksheet {_NS}><sheetData>'.encode()
    return entete + lignes + b"</sheetData></worksheet>"


def _cellule_texte(brut: bytes) -> bytes:
    """Cellule A1 portant un texte « en ligne », dans sa forme XML exacte."""
    return b'<row r="1"><c r="A1" t="inlineStr"><is><t>' + brut + b"</t></is></c></row>"


def _chaines_xml(textes: list[str]) -> str:
    entrees = "".join(f"<si><t>{t}</t></si>" for t in textes)
    return f'<?xml version="1.0"?><sst {_NS}>{entrees}</sst>'


def _ecrire_classeur(chemin, lignes, textes=None, onglet_actif=0, lignes2=""):
    """Écrit une archive xlsx minimale, entièrement maîtrisée.

    Passer par openpyxl rendrait les tests dépendants de sa façon d'échapper les
    caractères, qui varie selon les versions et la présence de lxml. Ici la
    forme exacte du XML est choisie par le test.
    """
    with zipfile.ZipFile(chemin, "w") as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("_rels/.rels", _RELS)
        archive.writestr("xl/workbook.xml", _workbook_xml(onglet_actif))
        archive.writestr("xl/_rels/workbook.xml.rels", _WORKBOOK_RELS)
        if textes is not None:
            archive.writestr("xl/sharedStrings.xml", _chaines_xml(textes))
        archive.writestr("xl/worksheets/sheet1.xml", _feuille_xml(lignes))
        archive.writestr("xl/worksheets/sheet2.xml", _feuille_xml(lignes2))


@pytest.fixture
def classeur_excel(tmp_path):
    """Classeur au format « Excel » : textes dans la table partagée."""

    def _creer(nom="excel.xlsx", onglet_actif=0, feuille2=""):
        textes = ["ID erreur", "Statut", "13001_1", "13001_2", "à revoir"]
        lignes = (
            '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c></row>'
            '<row r="2"><c r="A2" t="s"><v>2</v></c><c r="B2" t="s"><v>4</v></c></row>'
            '<row r="3"><c r="A3" t="s"><v>3</v></c><c r="B3"><v>42</v></c></row>'
        )
        chemin = tmp_path / nom
        _ecrire_classeur(chemin, lignes, textes, onglet_actif, feuille2)
        return str(chemin)

    return _creer


# ── Chaînes partagées (fichiers produits par Excel) ───────────────
def test_entetes_depuis_les_chaines_partagees(classeur_excel):
    assert lire_entetes(classeur_excel()) == {1: "ID erreur", 2: "Statut"}


def test_colonne_depuis_les_chaines_partagees(classeur_excel):
    assert list(lire_colonne(classeur_excel(), 1)) == [(2, "13001_1"), (3, "13001_2")]


def test_valeur_numerique_lue_telle_quelle(classeur_excel):
    assert list(lire_colonne(classeur_excel(), 2)) == [(2, "à revoir"), (3, "42")]


# ── Chaînes « en ligne » (fichiers produits par openpyxl) ─────────
def test_lecture_d_un_classeur_openpyxl(tmp_path):
    chemin = str(tmp_path / "openpyxl.xlsx")
    classeur = Workbook()
    feuille = classeur.active
    feuille.append(["ID erreur", "Statut"])
    feuille.append(["13001_1", "à revoir"])
    feuille.append(["13001_2", "à revoir"])
    classeur.save(chemin)

    assert lire_entetes(chemin) == {1: "ID erreur", 2: "Statut"}
    assert list(lire_colonne(chemin, 1)) == [(2, "13001_1"), (3, "13001_2")]


def test_entites_xml_restituees(tmp_path):
    """Accents et caractères réservés doivent revenir intacts.

    openpyxl échappe les accents en références numériques (« contr&#244;le ») :
    sans décodage, aucune colonne accentuée ne serait reconnue — cas courant
    dans les audits (« Date de contrôle », « État »).
    """
    chemin = str(tmp_path / "entites.xlsx")
    classeur = Workbook()
    classeur.active.append(["Date de contrôle", "État « à vérifier »", "A & B <x>"])
    classeur.save(chemin)

    assert lire_entetes(chemin) == {
        1: "Date de contrôle", 2: "État « à vérifier »", 3: "A & B <x>",
    }


@pytest.mark.parametrize(
    "brut,attendu",
    [
        ("contrôle".encode("utf-8"), "contrôle"),   # UTF-8 direct (lxml, Excel)
        (b"contr&#244;le", "contrôle"),             # référence décimale
        (b"contr&#xF4;le", "contrôle"),             # référence hexadécimale
        (b"contr&#XF4;le", "contrôle"),             # préfixe en majuscule
        (b"A &amp; B &lt;x&gt;", "A & B <x>"),      # entités nommées
    ],
)
def test_toutes_les_formes_d_echappement(tmp_path, brut, attendu):
    """Les accents s'écrivent de plusieurs façons selon le producteur du fichier.

    openpyxl les échappe en références numériques ou les laisse en UTF-8 selon
    sa version et la présence de lxml ; Excel a ses propres habitudes. Aucune de
    ces formes ne doit empêcher de reconnaître une colonne « Date de contrôle ».
    """
    chemin = tmp_path / "echappement.xlsx"
    _ecrire_classeur(chemin, _cellule_texte(brut))

    assert lire_entetes(str(chemin)) == {1: attendu}


def test_entetes_identiques_a_openpyxl(tmp_path):
    """Garde-fou : le chemin rapide et openpyxl doivent voir les mêmes colonnes."""
    from openpyxl import load_workbook

    from sna.services.colonnes_excel import entetes_depuis_feuille

    chemin = str(tmp_path / "comparaison.xlsx")
    classeur = Workbook()
    classeur.active.append([
        "Adresse", "ID erreur", "Date de contrôle", "Audit — Street View",
        "Nb BAL", "Commentaire « libre »",
    ])
    classeur.active.append(["rue", "13001_1", None, "à revoir", 3, "x & y"])
    classeur.save(chemin)

    reference = load_workbook(chemin, read_only=True)
    try:
        attendu = entetes_depuis_feuille(reference.active, 1)
    finally:
        reference.close()

    assert lire_entetes(chemin) == attendu


# ── Feuille active ────────────────────────────────────────────────
def test_onglet_actif_respecte(classeur_excel):
    """Le lecteur doit viser la même feuille qu'openpyxl : celle de activeTab."""
    lignes = '<row r="1"><c r="A1" t="inlineStr"><is><t>Autre feuille</t></is></c></row>'
    chemin = classeur_excel(onglet_actif=1, feuille2=lignes)

    assert lire_entetes(chemin) == {1: "Autre feuille"}


def test_onglet_actif_hors_limites_retombe_sur_le_premier(classeur_excel):
    assert lire_entetes(classeur_excel(onglet_actif=7)) == {1: "ID erreur", 2: "Statut"}


# ── Découpage en blocs ────────────────────────────────────────────
def test_resultat_identique_quelle_que_soit_la_taille_des_blocs(monkeypatch, tmp_path):
    """Le découpage du flux ne doit jamais scinder une cellule."""
    chemin = str(tmp_path / "gros.xlsx")
    classeur = Workbook()
    feuille = classeur.active
    feuille.append(["ID erreur"])
    for i in range(2, 400):
        feuille.append([f"13001_{i}"])
    classeur.save(chemin)

    complet = list(lire_colonne(chemin, 1))
    monkeypatch.setattr(lecture_rapide, "TAILLE_BLOC", 64)
    morcele = list(lire_colonne(chemin, 1))

    assert morcele == complet
    assert len(complet) == 398


# ── Refus explicites (déclenchent le repli openpyxl) ──────────────
def test_cellule_sans_reference_refusee(classeur_excel, tmp_path):
    chemin = str(tmp_path / "sans_ref.xlsx")
    with zipfile.ZipFile(classeur_excel()) as source, zipfile.ZipFile(chemin, "w") as cible:
        for nom in source.namelist():
            contenu = source.read(nom).decode("utf-8")
            if nom.endswith("sheet1.xml"):
                contenu = contenu.replace('<c r="A2" t="s">', "<c t=\"s\">")
            cible.writestr(nom, contenu)

    with pytest.raises(LectureImpossible):
        lire_entetes(chemin)


def test_fichier_illisible_refuse(tmp_path):
    chemin = tmp_path / "pas_un_xlsx.xlsx"
    chemin.write_text("ceci n'est pas une archive", encoding="utf-8")

    with pytest.raises(LectureImpossible):
        lire_entetes(str(chemin))


def test_archive_sans_classeur_refusee(tmp_path):
    chemin = str(tmp_path / "vide.xlsx")
    with zipfile.ZipFile(chemin, "w") as archive:
        archive.writestr("quelconque.txt", "x")

    with pytest.raises(LectureImpossible):
        lire_entetes(chemin)
