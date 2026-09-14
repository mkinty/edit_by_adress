"""Orchestration d'un lot de communes, une fois les codes INSEE résolus."""

from openpyxl import Workbook, load_workbook

from sna.models.base import Statut
from sna.models.config_adresses import ConfigAdresses
from sna.services.adresses import SaisieAdresse
from sna.services.traitement import traiter_communes

CONFIG = ConfigAdresses(colonne_adresse="Adresse", colonne_id="ID erreur")


def _s(adresse, valeur):
    return SaisieAdresse(adresse, valeur)


def _preparer_commune(racine, commune, adresses, valeurs_id=None):
    """Crée ``<racine>/Dep<XX>/<insee>/audit_<insee>.xlsx`` avec une ligne par adresse."""
    dossier = racine / f"Dep{commune[:2]}" / commune
    dossier.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Adresse"
    ws["C1"] = "ID erreur"
    valeurs_id = valeurs_id or ["x"] * len(adresses)
    for i, (adresse, valeur) in enumerate(zip(adresses, valeurs_id), start=2):
        ws[f"A{i}"] = adresse
        ws[f"C{i}"] = valeur
    chemin = dossier / f"audit_{commune}.xlsx"
    wb.save(chemin)
    return chemin


def test_traite_une_commune_par_groupe(tmp_path):
    _preparer_commune(tmp_path, "13001", ["12 rue A", "20 rue B", "5 rue C"])
    _preparer_commune(tmp_path, "75020", ["8 avenue D"])

    resultats = traiter_communes(
        {
            "13001": [_s("12 rue A", "13001_1"), _s("5 rue C", "13001_3")],
            "75020": [_s("8 avenue D", "75020_5")],
        },
        CONFIG, str(tmp_path),
    )

    assert [r.commune for r in resultats] == ["13001", "75020"]
    assert resultats[0].nb_modifications == 2
    assert resultats[1].nb_modifications == 1
    assert all(r.statut is Statut.SUCCES for r in resultats)


def test_adresses_absentes_du_fichier_sont_signalees(tmp_path):
    _preparer_commune(tmp_path, "13001", ["12 rue A"])

    resultat = traiter_communes(
        {"13001": [_s("12 rue A", "13001_1"), _s("adresse fantôme", "13001_42")]},
        CONFIG, str(tmp_path),
    )[0]

    assert resultat.statut is Statut.SUCCES
    assert resultat.adresses_absentes == ["adresse fantôme"]
    assert resultat.nb_adresses_demandees == 2
    assert resultat.resume_adresses() == "1/2 adresses trouvées"


def test_commune_sans_fichier(tmp_path):
    resultat = traiter_communes(
        {"99999": [_s("nulle part", "99999_1")]}, CONFIG, str(tmp_path)
    )[0]

    assert resultat.statut is Statut.INTROUVABLE
    assert resultat.nb_adresses_demandees == 1


def test_colonne_absente_donne_une_erreur_isolee(tmp_path):
    _preparer_commune(tmp_path, "13001", ["12 rue A"])
    _preparer_commune(tmp_path, "75020", ["8 avenue D"])
    config = ConfigAdresses(colonne_adresse="Adresse", colonne_id="Colonne absente")

    resultats = traiter_communes(
        {
            "13001": [_s("12 rue A", "13001_1")],
            "75020": [_s("8 avenue D", "75020_1")],
        },
        config, str(tmp_path),
    )

    assert all(r.statut is Statut.ERREUR for r in resultats)
    assert "introuvable" in resultats[0].message


def test_simulation_n_ecrit_rien(tmp_path):
    fichier = _preparer_commune(tmp_path, "13001", ["12 rue A"])

    resultat = traiter_communes(
        {"13001": [_s("12 rue A", "13001_1")]}, CONFIG, str(tmp_path), simuler=True
    )[0]

    assert resultat.nb_modifications == 1
    assert load_workbook(fichier).active["C2"].value == "x"


def test_callbacks_sont_appeles(tmp_path):
    _preparer_commune(tmp_path, "13001", ["12 rue A"])
    progressions, resultats = [], []

    traiter_communes(
        {"13001": [_s("12 rue A", "13001_1")]}, CONFIG, str(tmp_path),
        on_progression=lambda f, t, lib: progressions.append((f, t)),
        on_resultat=resultats.append,
    )

    assert (0, 1) in progressions and (1, 1) in progressions
    assert len(resultats) == 1 and resultats[0].commune == "13001"


def test_les_etapes_sont_transmises_avec_la_commune(tmp_path):
    """Le journal doit s'alimenter pendant le traitement, pas seulement après."""
    _preparer_commune(tmp_path, "13001", ["12 rue A"])
    etapes = []

    traiter_communes(
        {"13001": [_s("12 rue A", "13001_1")]}, CONFIG, str(tmp_path),
        on_etape=lambda commune, message: etapes.append((commune, message)),
    )

    assert etapes[0] == ("13001", "recherche du fichier audit…")
    assert any("audit_13001.xlsx" in message for _, message in etapes)
    assert all(commune == "13001" for commune, _ in etapes)


def test_duree_mesuree_par_commune(tmp_path):
    _preparer_commune(tmp_path, "13001", ["12 rue A"])

    resultat = traiter_communes(
        {"13001": [_s("12 rue A", "13001_1")]}, CONFIG, str(tmp_path)
    )[0]

    assert resultat.duree > 0
    assert resultat.resume_duree().endswith(" s")


def test_progression_annoncee_avec_le_nombre_d_adresses(tmp_path):
    _preparer_commune(tmp_path, "13001", ["12 rue A", "20 rue B"])
    libelles = []

    traiter_communes(
        {"13001": [_s("12 rue A", "13001_1"), _s("20 rue B", "13001_2")]},
        CONFIG, str(tmp_path),
        on_progression=lambda fait, total, libelle: libelles.append(libelle),
    )

    assert any("2 adresse(s)" in libelle for libelle in libelles)
